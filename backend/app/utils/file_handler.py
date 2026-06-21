"""文件处理工具：PDF/DOCX/MD/TXT 文本提取。"""
import io
from PyPDF2 import PdfReader
from docx import Document


def extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def extract_text_from_md_or_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "gbk", "utf-16"):
        try:
            return file_bytes.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore").strip()


def extract_text(filename: str, file_bytes: bytes) -> tuple[str, str]:
    """根据扩展名分发解析；返回 (text, format)。"""
    name = filename.lower()
    if name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes), "pdf"
    if name.endswith(".docx"):
        return extract_text_from_docx(file_bytes), "docx"
    if name.endswith(".md") or name.endswith(".markdown"):
        return extract_text_from_md_or_txt(file_bytes), "md"
    if name.endswith(".txt"):
        return extract_text_from_md_or_txt(file_bytes), "txt"
    raise ValueError(f"不支持的文件格式: {filename}")
